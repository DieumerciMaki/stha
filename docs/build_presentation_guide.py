"""Preparation for the supervisor meeting, based on implemented behavior."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=section.bottom_margin=Cm(1.8)
section.left_margin=section.right_margin=Cm(2)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(24)
doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(11)
header=section.header.paragraphs[0];header.text='Déchets urbains    Préparation de la présentation chez l’encadreur'
for run in header.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor(0,0,0)
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('Version locale    9 octobre 2026    Page ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8)

def p(text):return doc.add_paragraph(text)
step_number=0
def h(text):
    global step_number
    step_number=0
    return doc.add_heading(text,2)
def page(title):
    paragraph=doc.add_heading(title,1);paragraph.paragraph_format.page_break_before=True
def oral(text):
    paragraph=doc.add_paragraph();paragraph.add_run('À dire pendant la présentation. ').bold=True;paragraph.add_run('« '+text+' »')
def step(text):
    global step_number
    step_number+=1
    paragraph=doc.add_paragraph(str(step_number)+'.  '+text)
    paragraph.paragraph_format.left_indent=Cm(.55)
    paragraph.paragraph_format.first_line_indent=Cm(-.55)
def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    for i,row in enumerate(t.rows):
        cant_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant_split)
        for cell,width in zip(row.cells,widths):
            cell.width=Cm(width);cell.vertical_alignment=1
            props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                element=OxmlElement('w:'+side);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');borders.append(element)
            props.append(borders);margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                element=OxmlElement('w:'+side);element.set(qn('w:w'),'95');element.set(qn('w:type'),'dxa');margins.append(element)
            props.append(margins)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E9E9EC');props.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before=Pt(2);paragraph.paragraph_format.space_after=Pt(2)
                for run in paragraph.runs:run.font.size=Pt(9.2);run.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)
    p('')
def figure(path,width,caption):
    if path.exists():
        paragraph=doc.add_paragraph();paragraph.alignment=1
        paragraph.add_run().add_picture(str(path),width=Cm(width))
        paragraph.paragraph_format.space_after=Pt(5)
        paragraph.paragraph_format.keep_with_next=True
        caption_p=p(caption)
        for run in caption_p.runs:run.font.size=Pt(8.5)

evidence=json.loads((ROOT/'docs/qa/test-image-live.json').read_text(encoding='utf-8'))
result=evidence['result']
record_url='http://127.0.0.1:8001/#observations/'+evidence['observation_id']

doc.add_paragraph('Comprendre et présenter l’application de segmentation des déchets urbains',style='Title')
doc.add_paragraph('Guide de lecture pour une présentation de l’avancement chez l’encadreur',style='Subtitle')
p('Ce document explique ce que fait chaque partie de l’application, ce qui se passe lorsqu’une photo est analysée et comment présenter le prototype sans confondre fonctionnement technique et précision scientifique. Il s’adresse au porteur du travail de fin d’étude, pour préparer son explication et sa démonstration.')
p('La version locale fonctionne : une photographie est importée, analysée avec YOLO, affichée sous forme de masques puis conservée avec ses résultats. Le travail à présenter est un prototype opérationnel de segmentation 2D. L’adaptation et l’évaluation sur un corpus annoté de Goma restent à réaliser.')
h('Le sujet retenu')
p('Le logiciel concerne les déchets urbains à partir d’images. Il reprend les objectifs du Word de REDACTION : détecter et segmenter les zones de déchets, associer une localisation lorsqu’elle est disponible et proposer une interface de consultation pour les gestionnaires. Le périmètre accepté est la segmentation 2D avec YOLO.')
p('Un objet détecté possède une classe, un score de confiance et un masque. Un masque décrit les pixels attribués à cet objet. La réunion des masques donne les pixels reconnus comme déchets dans toute la photo. Elle ne garantit pas le contour complet d’un dépôt.')
h('Les écrans à connaître')
table(['Écran','Rôle'],[
('Tableau de bord','Vue des observations, de leurs dernières analyses et des vérifications en attente.'),
('Observations','Registre des photographies et accès à leur dossier.'),
('Analyses','Historique des traitements et comparaison des résultats enregistrés.'),
('Carte','Affichage des observations disposant de coordonnées.'),
('Vérifications','Examen humain des résultats avec décision et commentaire.'),
('Modèle YOLO','Provenance, classes, version et limites du modèle installé.'),
('Paramètres','Configuration du projet, mot de passe et création de comptes.')],[4,13])
oral('Je vais présenter le parcours complet d’une photographie, puis montrer comment les résultats sont conservés et vérifiés. Le logiciel fonctionne déjà, mais je ne présente pas encore une précision mesurée sur les déchets de Goma.')

page('Démarrage et tableau de bord')
h('Ouvrir le logiciel')
p('Dans le dossier PROJECT, double-cliquer sur LANCER.cmd et garder la console ouverte. L’application est disponible à http://127.0.0.1:8001. Si le serveur fonctionne déjà, ouvrir directement cette adresse. Se connecter avec son compte ; la création de l’administrateur concerne uniquement le premier accès à une base neuve. Ctrl+C dans la console arrête le serveur web.')
p('La barre latérale donne accès aux sept écrans. La page active est mise en évidence. Sur ordinateur, le bouton en haut réduit la navigation aux icônes. Sur mobile, le bouton de menu l’ouvre. En bas, le compte permet d’accéder aux paramètres et de se déconnecter. Les icônes proviennent de Lucide.')
figure(ROOT/'docs/qa/sidebar-map-desktop.png',16,'Figure 1 Navigation et tableau de bord après correction de la carte. Capture de contrôle avec des données d’interface simulées uniquement dans le navigateur et un registre vide ; elle ne représente pas une collecte réelle.')
h('Lire les quatre indicateurs')
p('Observations compte les photos retenues par les filtres. Analyses terminées compte les observations dont la dernière analyse est terminée. À vérifier compte les dernières analyses terminées sans décision définitive. Observations localisées compte les photos disposant de coordonnées. Ces indicateurs ne comptent donc pas tous les traitements historiques d’une même photo.')
p('Le tableau récent ouvre les dossiers. Dernière analyse donne un aperçu du traitement le plus récemment lancé parmi les observations filtrées. La carte et la liste des vérifications complètent cette vue. La période filtre la date d’ajout de la photo ; le quartier provient des valeurs saisies. Le bouton d’actualisation recharge les données.')
oral('Cette page sert à suivre le registre. Les chiffres sont calculés à partir des données enregistrées ; ils ne sont pas des valeurs copiées depuis la maquette.')

page('Observations et importation des photographies')
h('Le rôle du registre')
p('Une observation associe une photo à un titre, à des informations de contexte et, éventuellement, à des coordonnées. Elle existe indépendamment de l’analyse. Une même photo peut faire l’objet de plusieurs traitements, par exemple avec des seuils de confiance différents.')
p('La page Observations présente une miniature, le titre, le quartier, la date d’ajout et l’état de la dernière analyse. La recherche porte sur le titre, le quartier et le lieu. Le filtre Analyse distingue les observations à analyser, en attente, en cours, terminées ou en échec. Le registre complet utilise une pagination de dix lignes. Ouvrir donne accès au dossier.')
h('Importer une image')
step('Cliquer sur Ajouter une observation et choisir une photographie. Les formats autorisés sont JPEG, PNG et WebP, avec un maximum de 15 Mo et de 20 millions de pixels.')
step('Saisir un titre obligatoire. Renseigner le quartier, le lieu ou repère, la date de capture et les notes seulement si ces informations sont connues.')
step('Renseigner ensemble latitude et longitude, ou laisser les deux vides. Le logiciel peut récupérer le GPS EXIF si la photo le contient et qu’aucune coordonnée manuelle n’est saisie.')
step('Conserver Lancer l’analyse après l’import coché pour analyser automatiquement, ou le décocher pour enregistrer la photo avant un traitement ultérieur.')
step('Cliquer sur Enregistrer l’observation. Le dossier s’ouvre ; si une analyse est demandée, elle est mise en file puis exécutée en arrière-plan.')
h('Comprendre les dates et les fichiers')
p('La date de capture est fournie par l’utilisateur ; elle est facultative. La date d’ajout est créée lors de l’enregistrement. Les dates affichées ne doivent pas être prises pour la preuve du moment exact de prise de vue. Le fichier importé est conservé et une version JPEG normalisée est produite pour le traitement, notamment après correction de l’orientation EXIF.')
p('Utiliser une image exemple propose trois photos externes attribuées à leur source. Le titre et les notes en indiquent la provenance. Ces exemples permettent une démonstration ; ils ne constituent pas des observations collectées à Goma.')
h('Exporter le registre')
p('Exporter CSV fournit le registre sous forme tabulaire, utilisable dans Excel. Cet export concerne l’ensemble du registre, même lorsqu’un filtre d’écran est actif. Le CSV facilite le suivi ; il ne contient pas les masques de pixels.')
oral('Je commence par enregistrer une observation avec sa photo et son contexte. La localisation est facultative, et je ne saisis pas un quartier ou des coordonnées que je ne connais pas.')

page('Le dossier et les vues de segmentation')
figure(ROOT/'docs/qa/test-image-interface.png',13.2,'Figure 2 Analyse réelle de la photographie exemple conservée dans le registre. La capture montre un compte technique de test, désormais désactivé. Le résultat reste accessible depuis le compte de l’utilisateur.')
h('Comparer les trois vues')
p('Original affiche la photo normalisée utilisée pour l’analyse. Segmentation superpose des couleurs et contours aux zones reconnues. Masque affiche en blanc l’union des zones segmentées et en noir le fond. Une couleur sert à distinguer une classe ; elle n’indique pas un niveau de danger. Le zoom permet d’inspecter une zone plus précisément.')
p('Le dossier affiche les dimensions de la photo, les notes et la localisation connue. Résultat affiché permet de choisir une analyse historique. Le bouton Lancer l’analyse crée un nouveau traitement au seuil choisi, sans effacer les anciens résultats. Retour au registre permet de reprendre la consultation.')
oral('Je compare les zones colorées à la photo originale. Ces zones sont des prédictions à examiner ; leur couleur ne garantit pas que le matériau est correctement reconnu.')

page('Lire et interpréter les résultats de la photo')
h('Le test réellement effectué')
p('Une photographie externe montrant des déchets parmi de la végétation a été importée par le formulaire de l’application. L’analyse a utilisé le modèle installé, sans résultats simulés. L’affichage, le masque, l’export JSON et la présence dans l’historique ont été contrôlés. Aucune coordonnée géographique n’a été inventée.')
p('Photo de '+str(result['width'])+' × '+str(result['height'])+' pixels, seuil de confiance de '+str(result['confidence_threshold']).replace('.',',')+' et calcul sur CPU. La décision reste À vérifier.')
table(['Mesure','Résultat du test'],[
('Instances détectées',str(result['instance_count'])),
('Couverture de l’image',str(result['coverage_percent']).replace('.',',')+' %'),
('Durée mesurée',str(result['duration_seconds']).replace('.',',')+' secondes'),
],[7,10])
h('Les matériaux prédits')
table(['Instance','Classe prédite','Confiance'],[(i+1,item['label'],f"{100*item['confidence']:.1f}".replace('.',',')+' %') for i,item in enumerate(result['instances'])],[2,9,6])
p('Ces classes sont les sorties du modèle. Elles ne sont pas des annotations de référence. L’observation visuelle montre aussi des zones de déchets non segmentées ; les matériaux reconnus doivent être contrôlés. Le test prouve que le parcours fonctionne, sans établir la précision générale du système.')
h('Le sens des nombres')
p('Une instance est un objet individualisé par le modèle. Ce n’est ni une masse de déchets ni un nombre de dépôts. La confiance est un score associé à une prédiction. Un score de 92,2 % ne signifie pas que le système possède une précision globale de 92,2 %.')
p('La couverture est le nombre de pixels blancs du masque global divisé par le nombre total de pixels, puis multiplié par 100. Les pixels communs à plusieurs masques ne sont comptés qu’une fois. Cette proportion ne mesure ni une surface au sol, ni un poids, ni un volume. La durée affichée est celle du traitement mesuré après chargement du modèle ; elle ne couvre pas toute l’attente entre le clic et l’affichage.')
oral('Sur cette image, le modèle a produit cinq instances et une couverture d’environ 3,58 %. Ce résultat montre le fonctionnement de la chaîne. Je garde la décision À vérifier parce que la qualité de la segmentation doit encore être examinée.')

page('Historique des analyses et vérifications')
h('La page Analyses')
p('Cette page contient les traitements enregistrés, et non seulement les photos. Chaque ligne indique l’observation, la date de lancement, l’état, le seuil, le nombre d’instances et la couverture lorsque le traitement est terminé. La recherche et les filtres facilitent la consultation. Ouvrir sélectionne précisément le résultat correspondant dans le dossier.')
table(['État','Signification'],[
('En attente','Traitement enregistré dans la file, pas encore exécuté.'),
('En cours','Le moteur effectue l’analyse.'),
('Terminée','Calcul achevé et résultats enregistrés. Cela ne signifie pas résultat validé.'),
('Échec','Le traitement a rencontré une erreur, consultable dans le dossier.')],[4,13])
p('Un seuil plus faible peut conserver davantage de prédictions, mais aussi davantage de faux positifs. Un seuil plus élevé peut supprimer des détections, y compris certains déchets réels. Le choix du seuil doit être justifié par une évaluation, pas uniquement par le nombre d’objets affichés.')
h('La page Vérifications')
p('Elle regroupe les résultats terminés selon leur décision : À vérifier, Acceptés, Rejetés ou Tous les résultats. Le bouton Examiner ouvre le dossier. L’administrateur ou le vérificateur compare les vues, choisit une décision et peut ajouter un commentaire.')
p('Accepter après examen signifie que l’utilisateur juge le résultat utilisable dans le contexte de l’observation. Rejeter indique qu’il ne le juge pas utilisable. À vérifier conserve une incertitude. Chaque décision est ajoutée au journal avec son auteur et sa date ; les décisions antérieures restent conservées.')
p('Un agent peut importer et analyser, mais ne peut pas enregistrer une vérification. La validation humaine ne modifie pas les contours ou les classes produits par YOLO. La version actuelle ne possède pas encore d’éditeur manuel de masques.')
h('Les exports du dossier')
p('Masque PNG fournit le masque global. Résultat JSON fournit les informations du traitement, dont les classes, scores, pixels, contours et paramètres. Exporter le dossier regroupe les informations de l’observation, ses analyses et leurs vérifications dans un fichier JSON. Ces fichiers permettent une consultation ou un traitement ultérieur.')
oral('Je distingue l’état technique du traitement de la décision humaine. Une analyse terminée peut encore être fausse. C’est pour cette raison que le logiciel possède un journal de vérification séparé.')

page('La carte et la localisation')
h('Ce que la carte représente')
p('La carte affiche les observations qui possèdent une latitude et une longitude. Le fond est fourni par OpenStreetMap et l’affichage interactif par Leaflet. Sans observations géolocalisées, la carte reste centrée sur Goma avec un message indiquant qu’aucune observation n’est localisée.')
p('Les coordonnées peuvent provenir de la saisie de l’utilisateur ou des métadonnées GPS EXIF de la photo. Si des coordonnées sont saisies, elles ont priorité sur le GPS EXIF. Une photo qui ne contient pas de GPS et dont le lieu n’est pas saisi reste sans position.')
h('Deux sens différents du mot localisation')
p('Dans l’image, YOLO localise un objet par ses pixels, son masque et sa boîte englobante. Sur le territoire, la carte localise une observation par ses coordonnées. Le modèle n’obtient pas une adresse ou un quartier à partir du seul masque. Le point sur la carte correspond à la position renseignée ou enregistrée par la photo, pas à une mesure géographique de chaque déchet.')
h('Utiliser la carte')
step('Ouvrir Carte dans la barre latérale. Appliquer une période ou un quartier si nécessaire.')
step('Utiliser les boutons plus et moins pour changer de niveau de zoom. Les observations disponibles déterminent le cadrage.')
step('Cliquer sur un point pour lire son titre et son contexte, puis sur Ouvrir l’observation pour examiner la photo.')
step('Utiliser Exporter GeoJSON pour récupérer les observations géolocalisées. Comme le CSV, cet export porte sur l’ensemble du registre, pas seulement sur les points filtrés à l’écran.')
p('La couleur du point distingue une dernière analyse terminée d’un autre état. Elle ne représente pas la gravité de la pollution. Le GeoJSON utilise des coordonnées dans l’ordre longitude puis latitude, selon le format géographique.')
h('Connexion et fonctionnement local')
p('Le registre et le calcul YOLO fonctionnent localement une fois les dépendances et le modèle installés. Le fond de carte nécessite une connexion Internet. Le blocage initial des tuiles a été corrigé par une politique de référence explicite sur les requêtes Leaflet. Un contrôle navigateur a obtenu 33 réponses HTTP 200 et vérifié l’affichage de Goma.')
oral('La carte permet de retrouver les observations dont le lieu est connu. Dans la photo exemple, je laisse la localisation vide, parce que je ne dispose pas de ses coordonnées réelles.')

page('Modèle YOLO paramètres et comptes')
h('La page Modèle YOLO')
p('Elle explique quel modèle est installé : architecture, jeu d’entraînement déclaré, classes reconnues, taille des poids, source, révision et empreinte SHA256. Le modèle courant est YOLO11s de segmentation, fourni par CatSat et entraîné sur TACO selon sa fiche. Ses cinq classes sont plastique, métal, verre, papier et autres.')
p('Le traitement est configuré avec une taille d’entrée de 960 pixels et un calcul CPU. Les masques sont produits à la résolution de la photo. Prêt indique que les fichiers du modèle sont configurés ; lors du premier calcul, le moteur charge les poids et contrôle leur empreinte ainsi que leur compatibilité.')
p('Les scores de masque publiés par le fournisseur sur TACO sont 27,2 % de mAP à IoU 0,50 et 18,8 % de mAP sur les seuils IoU de 0,50 à 0,95. Ils sont affichés comme résultats externes. Ils ne décrivent pas une évaluation indépendante du projet sur Goma. L’interface indique qu’aucune évaluation de terrain n’est encore enregistrée.')
h('La page Paramètres')
p('L’administrateur peut modifier le nom du projet, la ville et le seuil de confiance par défaut. Modifier le seuil par défaut affecte les futures analyses ; les anciens résultats conservent leur seuil. Modifier la ville affichée ne déplace pas automatiquement les coordonnées ni le centrage initial de la carte.')
p('Chaque utilisateur peut modifier son mot de passe en fournissant le mot de passe actuel. Un administrateur peut consulter les comptes et en créer avec un nom, un identifiant, un mot de passe initial et un rôle. La version actuelle ne présente pas une interface complète de suppression ou de modification des rôles des comptes existants.')
table(['Action','Administrateur','Agent','Vérificateur'],[
('Consulter et exporter','Oui','Oui','Oui'),
('Importer et analyser','Oui','Oui','Non'),
('Enregistrer une vérification','Oui','Non','Oui'),
('Configurer le projet et créer des comptes','Oui','Non','Non'),
('Modifier son mot de passe','Oui','Oui','Oui')],[8,3,3,3])
h('La séparation des droits')
p('Les droits sont contrôlés par le serveur, pas seulement par les boutons de l’interface. Les données et les fichiers du registre nécessitent une session authentifiée. Un compte technique agent a servi au test de la photo ; il a ensuite été désactivé. Les résultats restent dans le registre pour la présentation.')
oral('Le modèle est externe et sa provenance est visible. Mon travail présenté ici concerne l’intégration, le traitement, la traçabilité et l’interface. L’entraînement et l’évaluation sur les données de Goma constituent la prochaine étape.')

page('Architecture et conservation des données')
h('Le trajet d’une demande')
p('Le navigateur affiche React et envoie une demande HTTP à FastAPI. FastAPI vérifie la session, les droits et les données. PostgreSQL conserve les informations structurées. Pour une analyse, le serveur enregistre une tâche ; le moteur utilise YOLO et OpenCV, écrit les fichiers de résultat puis met à jour l’état. Le navigateur recharge les données pendant le traitement pour afficher son évolution.')
table(['Technologie','Rôle dans cette version'],[
('React TypeScript et Vite','Interface, formulaires, navigation et construction du code destiné au navigateur.'),
('CSS Tailwind et Lucide','Mise en page, styles et icônes SVG. La barre latérale adapte une structure inspirée de shadcn/ui.'),
('FastAPI et Python','API, validation des imports, sessions, droits, tâches et exports.'),
('Ultralytics et PyTorch','Chargement et exécution du modèle YOLO de segmentation.'),
('OpenCV NumPy et Pillow','Lecture, orientation, manipulation des pixels, union des masques, contours et images de résultat.'),
('PostgreSQL SQLAlchemy et Alembic','Stockage structuré, accès aux données et évolution du schéma.'),
('Leaflet et OpenStreetMap','Carte interactive et fond cartographique distant.')],[6,11])
h('Pourquoi la base et les fichiers sont séparés')
p('PostgreSQL conserve les utilisateurs, sessions, observations, analyses, vérifications et paramètres. Les photos originales, copies normalisées, superpositions et masques sont conservés dans les dossiers de données. Les poids du modèle sont dans models. Cela évite d’enregistrer chaque grande image directement dans les tables.')
p('La base propre au projet fonctionne sur le port local 5433 ; le serveur web utilise 8001. Les fichiers backend/app/main.py, segmentation.py et database.py portent respectivement l’API, le moteur et les structures de données. frontend/src contient les écrans. scripts contient notamment les outils de préparation des données, entraînement, évaluation et sauvegarde.')
h('Sécurité et traçabilité')
p('Les mots de passe sont hachés avec scrypt et un sel. La session utilise un cookie HttpOnly et un jeton dont l’empreinte est conservée en base. Des contrôles limitent les tentatives de connexion et refusent les mutations depuis une origine non autorisée. Les empreintes SHA256 des images et du modèle permettent de reconnaître les fichiers utilisés ; elles ne prouvent pas l’exactitude des prédictions.')
p('En local, les traitements utilisent un exécuteur en arrière-plan. Celery, Redis et Docker sont fournis comme configuration d’évolution ; ils ne sont pas les services utilisés dans la démonstration Windows et n’ont pas été validés sur ce PC. CVAT est un outil externe d’annotation, pas une page intégrée à l’application.')

page('Comment une photographie est analysée')
h('La chaîne de calcul réelle')
step('FastAPI contrôle le type, la taille et la validité de la photo. Pillow corrige son orientation et produit une image JPEG normalisée. L’original et les empreintes sont conservés.')
step('La demande d’analyse enregistre le seuil, l’auteur, la date et l’état En attente. Le traitement en arrière-plan passe ensuite à En cours.')
step('Le moteur charge une fois les poids YOLO, vérifie leur SHA256 et leur tâche de segmentation. OpenCV lit la photo à partir du fichier normalisé.')
step('Ultralytics exécute YOLO sur le CPU avec imgsz égal à 960 et retina_masks activé. Pour chaque instance retenue, il produit une classe, un score, une boîte et un masque.')
step('NumPy et OpenCV réunissent les pixels des masques, comptent leur couverture, extraient les contours et composent la superposition colorée. Les trous des masques sont conservés dans le masque raster global.')
step('Le moteur écrit overlay.jpg, mask.png et result.json. L’analyse devient Terminée ; en cas d’erreur, elle devient Échec avec un message. Le navigateur affiche les fichiers et les informations enregistrées.')
h('Ce que le modèle a appris')
p('Les poids sont les paramètres appris pendant l’entraînement. Le modèle installé a déjà appris à reconnaître certaines apparences de déchets sur un corpus externe. Le lancer sur une nouvelle photo est une inférence. Importer une photo dans le registre ou accepter un résultat ne réentraîne pas automatiquement le modèle.')
h('Comment mesurer la qualité pour le mémoire')
p('Il faudra constituer un corpus de photos de Goma et des masques de référence annotés, par exemple avec CVAT. Séparer les données d’entraînement, de validation et de test ; des photos très proches d’un même site ne doivent pas se répartir artificiellement entre ces ensembles. Le test doit rester indépendant du choix des paramètres.')
p('La précision mesure la proportion de prédictions correctes ; le rappel mesure la proportion des objets de référence retrouvés. L’IoU compare le recouvrement entre un masque prédit et un masque de référence : intersection divisée par union. La mAP résume les performances selon les classes, seuils de confiance et critères de recouvrement. Elle nécessite des annotations de référence et ne se déduit pas du compteur d’instances.')
p('Les scripts de conversion COCO, contrôle des données, entraînement, enregistrement du modèle et évaluation sont présents. Ils préparent cette démarche ; aucun entraînement ni mesure scientifique sur Goma n’est annoncé comme effectué.')
oral('YOLO produit les prédictions. OpenCV prépare leur représentation et leurs mesures en pixels. La justesse de ces prédictions sera évaluée avec des annotations indépendantes de terrain.')

page('Déroulement de la démonstration chez l’encadreur')
p('Durée indicative : huit à dix minutes, à adapter au temps accordé. Commencer par le problème, montrer une photo et suivre son parcours. Garder les détails techniques pour l’explication et les questions.')
table(['Étape','Ce que je montre','Ce que j’explique'],[
('1','Tableau de bord','Le registre de photos et le suivi des analyses.'),
('2','Ajouter une observation','Le titre, le contexte, les coordonnées facultatives et l’import réel.'),
('3','Original puis Segmentation','Les zones produites par YOLO et leur comparaison avec la photo.'),
('4','Masque et tableau des instances','Les pixels retenus, matériaux prédits, confiance et couverture.'),
('5','Analyses puis Vérifications','L’historique technique et le contrôle humain séparé.'),
('6','Carte','Les points proviennent de coordonnées connues ; l’exemple n’est pas géolocalisé.'),
('7','Modèle YOLO','La provenance des poids, les classes et l’absence d’évaluation de terrain.'),
('8','Conclusion orale','Les résultats obtenus et le travail restant sur le corpus de Goma.')],[1.6,6.2,9.2])
h('Une introduction à apprendre')
oral('Mon travail porte sur la détection et la segmentation des déchets urbains à partir de photographies, avec YOLO et OpenCV. L’objectif de cette version est de permettre l’importation, l’analyse, la consultation et la vérification des résultats dans une application locale. Je vais montrer le parcours d’une image, puis expliquer les limites et les prochaines étapes expérimentales.')
h('Une conclusion à apprendre')
oral('Le parcours technique est opérationnel et un traitement réel a été vérifié. Il reste à constituer un corpus annoté de Goma, adapter le modèle et mesurer ses performances sur des photos indépendantes. Je souhaite valider avec vous les classes à retenir, le protocole d’annotation et les critères d’évaluation.')
h('Préparer le matériel')
p('Vérifier avant la séance que le serveur est actif, que le compte fonctionne et que le modèle est prêt. Ouvrir l’image déjà analysée, préparer une photo exemple pour l’import et garder les captures d’écran disponibles. Éviter une installation, un changement de poids ou une nouvelle fonctionnalité juste avant la démonstration.')
p('Résultat réel déjà conservé : '+record_url)
p('Si une nouvelle analyse est lente, expliquer que le premier chargement peut prendre davantage de temps, puis ouvrir le résultat enregistré. Si Internet manque, présenter les écrans locaux et expliquer la dépendance du fond de carte. Ne pas remplacer une erreur par des nombres inventés.')

page('Questions possibles et prochaines décisions')
h('Pourquoi YOLO et OpenCV')
p('YOLO fournit la reconnaissance et les masques d’instances. OpenCV manipule les images et les résultats, notamment les contours et la superposition. Les deux ont des rôles complémentaires. La méthode utilisée n’est pas une comparaison d’histogrammes pour détecter des changements de scènes vidéo.')
h('Est ce que le modèle reconnaît tous les déchets')
p('Non. Il possède cinq classes de matériaux et peut manquer des objets, confondre leurs matériaux ou réagir à des éléments du fond. La végétation, les occlusions, la petite taille des objets et les différences avec les images d’entraînement peuvent compliquer l’analyse. Leur effet précis sur Goma doit être mesuré.')
h('Est ce que le logiciel calcule le volume ou le danger')
p('Non. Il calcule une couverture en pixels. Sans données de profondeur, échelle physique et méthode validée, il ne fournit ni volume, ni masse, ni surface au sol. Le type de matériau ne suffit pas à déterminer automatiquement la dangerosité.')
h('Est ce que la vérification améliore automatiquement YOLO')
p('Non. La vérification conserve une décision et un commentaire. Pour réentraîner, il faut des annotations de masques, préparer les partitions, exécuter un entraînement et évaluer le nouveau modèle. La correction manuelle des contours et le mode détaillé par découpes ont été proposés, mais ne sont pas encore intégrés.')
h('Qu est ce qui est déjà prouvé')
p('Le fonctionnement local de l’import, du calcul réel, de l’affichage, des exports et de l’historique a été contrôlé. Six tests automatisés ont passé lors de la validation de la version locale. La carte et la navigation ont aussi été contrôlées dans le navigateur. Ce sont des validations techniques, pas des statistiques de reconnaissance sur un jeu de test de Goma.')
h('Les points à valider avec l’encadreur')
p('Confirmer le périmètre image 2D, les classes de déchets pertinentes, les lieux et conditions de collecte, la procédure d’annotation et les métriques à présenter. Le Word contient des passages de gabarit et une proposition topologique 3D à harmoniser avec ce périmètre. Leur correction rédactionnelle doit accompagner la méthode réellement mise en œuvre.')
p('L’application peut évoluer vers une analyse détaillée par découpes, une inspection objet par objet et un éditeur de corrections. La priorité scientifique reste une évaluation indépendante, afin de savoir si ces changements améliorent la précision ou le rappel et à quel coût de calcul.')
h('Sources et documents de référence')
p('Mémoire original dans REDACTION ; correspondance documentée dans PROJECT/docs/COHERENCE_MEMOIRE.md. Code relu : App.tsx, AnalysisDetail.tsx, forms.tsx, Sidebar.tsx, components.tsx ; main.py, auth.py, database.py et segmentation.py. Mesures de la photo : docs/qa/test-image-live.json.')
for text in ('Fiche du modèle : https://huggingface.co/CatSat/yolov11-litter-materials',
             'Documentation de segmentation : https://docs.ultralytics.com/tasks/segment/',
             'Guide complémentaire : PROJECT/docs/GUIDE_UTILISATION.docx ; validation : PROJECT/docs/VALIDATION_LOCALE.md.'):
    paragraph=p(text)
    for run in paragraph.runs:run.font.size=Pt(8.5)

doc.core_properties.title='Comprendre et présenter l’application de segmentation des déchets urbains'
doc.core_properties.subject='Préparation de la présentation de l’avancement chez l’encadreur'
doc.core_properties.author=''
doc.save(ROOT/'docs/PREPARATION_PRESENTATION_ENCADREUR.docx')
print(ROOT/'docs/PREPARATION_PRESENTATION_ENCADREUR.docx')

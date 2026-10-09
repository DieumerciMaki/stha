"""Create the user and technical guide with the bundled document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=section.bottom_margin=Cm(1.8)
section.left_margin=section.right_margin=Cm(2)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(8)
for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(14)
header=section.header.paragraphs[0];header.text='Déchets urbains    Guide d’utilisation et de fonctionnement'
for run in header.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor(0,0,0)
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('PROJECT    8 octobre 2026    Page ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8)

def p(text):doc.add_paragraph(text)
def h(text):doc.add_heading(text,level=2)
def page(title):
    heading=doc.add_heading(title,level=1);heading.paragraph_format.page_break_before=True
def step(text):doc.add_paragraph(text,style='List Number')
def code(text):
    text=text.replace('\n  --',' `\n  --')
    para=doc.add_paragraph();run=para.add_run(text);run.font.name='Consolas';run.font.size=Pt(8.5)
    para.paragraph_format.space_after=Pt(9)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    for cell,value in zip(t.rows[0].cells,headers):cell.text=value
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for i,row in enumerate(t.rows):
        for cell,width in zip(row.cells,widths):
            cell.width=Cm(width);cell.vertical_alignment=1
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF1');props.append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
                for run in para.runs:run.font.size=Pt(9);run.font.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

doc.add_paragraph('Guide du système de segmentation des déchets urbains',style='Title')
doc.add_paragraph('Application du travail de fin d’étude dans le dossier PROJECT',style='Subtitle')
p('Ce guide permet de lancer l’application, d’importer une photographie, d’examiner une segmentation YOLO et d’enregistrer une vérification. Il explique aussi la base de données, les fichiers du projet et la démarche nécessaire pour entraîner et évaluer le modèle sur des images de Goma.')
p('La version locale est installée sur ce PC. Le moteur utilise des poids réellement entraînés sur TACO et produit des masques de pixels. Son fonctionnement technique a été testé ; aucune précision scientifique sur Goma n’est encore mesurée.')
h('Démarrer et faire un premier essai')
step('Ouvrez PROJECT et double-cliquez sur LANCER.cmd. Gardez la fenêtre de lancement ouverte. Si le serveur est déjà actif, passez directement à l’adresse http://127.0.0.1:8001.')
step('Au premier accès, créez votre compte administrateur. Indiquez votre nom, choisissez un identifiant et un mot de passe d’au moins dix caractères. Aux accès suivants, utilisez ces informations pour vous connecter.')
step('Dans le tableau de bord, cliquez sur Ajouter une observation. Choisissez votre photo ou cliquez sur Utiliser une image exemple. Trois exemples externes sont fournis pour tester immédiatement le modèle.')
step('Saisissez un titre. Ajoutez les informations réellement connues : quartier, lieu, date de capture et notes. Les coordonnées sont facultatives ; latitude et longitude doivent être renseignées ensemble.')
step('Gardez Lancer l’analyse après l’import coché, puis enregistrez. Le traitement s’exécute sur le CPU en arrière-plan. Le premier chargement du modèle peut prendre plus longtemps que les traitements suivants.')
step('Après le traitement, comparez Original, Segmentation et Masque. Enregistrez une décision et un commentaire après examen. Téléchargez le masque PNG ou le dossier JSON si nécessaire.')
h('Arrêter')
p('Dans la fenêtre du serveur, utilisez Ctrl+C. Les observations restent conservées. PostgreSQL du projet peut rester actif sur le port 5433. Le lancement utilise exclusivement l’adresse locale 127.0.0.1.')

page('Comprendre le tableau de bord')
p('La disposition reprend la maquette fournie : navigation à gauche, indicateurs compacts, registre récent, aperçu d’analyse, carte et résultats en attente. Les nombres sont calculés depuis la base ; les chiffres de la maquette ne sont pas copiés.')
image=root/'docs/qa/dashboard-real.png'
if image.exists():
    doc.add_picture(str(image),width=Cm(17))
    para=doc.add_paragraph('Figure 1 Interface vérifiée avec une image exemple externe sans coordonnées de terrain.');para.runs[0].font.size=Pt(8.5)
h('Définition des indicateurs')
p('Observations compte les photographies correspondant aux filtres de période et de quartier. Analyses terminées compte les observations dont la dernière analyse est achevée. À vérifier compte les dernières analyses achevées sans décision définitive. Observations localisées compte les photographies ayant une paire de coordonnées.')
p('La période s’applique à la date d’ajout au registre. Les filtres concernent l’affichage. Les exports CSV et GeoJSON portent sur l’ensemble du registre, et pas uniquement sur les lignes visibles.')
p('Un registre neuf affiche zéro. Les images exemples ne sont ajoutées que lorsque vous les importez. Leur titre et leurs notes indiquent leur provenance externe ; elles restent sans coordonnées tant que leur lieu n’est pas connu.')

page('Importer consulter et vérifier')
h('Informations de la photographie')
p('Le logiciel accepte JPEG, PNG et WebP, jusqu’à 15 Mo et 20 millions de pixels. Il conserve le fichier importé, crée une image JPEG normalisée et applique l’orientation EXIF. Deux empreintes SHA256 permettent d’identifier le fichier source et l’image réellement analysée.')
p('Les coordonnées saisies sont prioritaires. Sans saisie, le logiciel utilise les coordonnées GPS EXIF lorsqu’elles existent. En leur absence, il laisse le lieu géographique vide. Le nom du quartier ne permet pas de calculer une position GPS.')
h('Examiner une analyse')
p('Ouvrez une observation depuis le registre, le dashboard, la carte ou l’historique. La vue Original affiche la photographie ; Segmentation affiche les masques colorés ; Masque affiche l’union binaire des pixels reconnus comme déchets. Le zoom permet d’examiner une région. Le sélecteur de résultat conserve l’accès aux analyses précédentes.')
p('Le seuil de confiance filtre les prédictions. Un seuil plus faible peut augmenter les détections et les faux positifs. Vous pouvez relancer une analyse avec un autre seuil ; chaque traitement conserve son résultat, son modèle et son horodatage. Aucune détection ne signifie pas nécessairement absence de déchets.')
h('Vérification et comptes')
table(['Rôle','Actions autorisées'],[('Agent','Importer, lancer une analyse, consulter et exporter.'),('Vérificateur','Consulter, exporter et enregistrer une décision.'),('Administrateur','Toutes les actions, paramètres et création de comptes.')],[4,13])
p('Après examen, choisissez Accepter après examen, Rejeter ou À vérifier. Une décision est ajoutée au journal avec votre nom et la date. Les anciennes décisions restent consultables. Une acceptation utilisateur ne constitue pas une mesure statistique de précision.')
h('Exports')
p('Le dossier d’une observation contient ses métadonnées, les analyses et les vérifications. Le résultat JSON contient les classes, scores, polygones et empreinte du modèle. Le masque PNG conserve l’union exacte des masques raster. CSV facilite la consultation du registre ; GeoJSON exporte seulement les points ayant des coordonnées.')

page('Fonctionnement de la segmentation')
h('Modèle installé')
p('Le modèle YOLO11s-seg provient de CatSat/yolov11-litter-materials sur Hugging Face. Il a été adapté à TACO en regroupant les classes par matériau : plastique, métal, verre, papier et autres. Ses poids sont installés dans models/waste-yolo11s-seg.pt. Le manifeste active.json enregistre la source, la révision et l’empreinte vérifiée avant chargement.')
h('Chaîne de traitement')
p('React envoie la photo à FastAPI. L’API valide le fichier, enregistre l’observation dans PostgreSQL et ajoute un traitement à la file. PyTorch et Ultralytics exécutent YOLO. OpenCV et NumPy composent les masques, tracent les contours et produisent les fichiers. L’API conserve les résultats, puis l’interface les affiche.')
p('YOLO réalise une segmentation d’instances. Le logiciel construit ensuite une union binaire déchets/fond. Les pixels communs à plusieurs instances ne sont comptés qu’une fois et les trous présents dans les masques raster sont préservés. Les polygones exportés servent à décrire les contours ; le PNG est la référence du masque global raster.')
code('Couverture (%) = 100 × pixels de l’union des masques / pixels de l’image')
p('Cette valeur décrit une image 2D. Elle ne donne ni des mètres carrés, ni un volume, ni le nombre de dépôts d’un quartier. Les coordonnées indiquent le point associé à la photographie, pas la position GPS précise de chaque pixel segmenté.')
h('Résultats publiés et résultats locaux')
table(['Mesure','Valeur publiée sur TACO'],[('mAP masque 50','27,2 %'),('mAP masque 50 à 95','18,8 %')],[9,8])
p('Ces deux valeurs sont déclarées par l’auteur du modèle externe. Elles ne sont pas des scores obtenus sur Goma. Le modèle peut confondre les matériaux, manquer de petits objets ou ne segmenter qu’une partie d’un amas. Une évaluation indépendante de terrain est nécessaire pour le mémoire.')
p('Un essai technique de l’application a produit cinq instances sur l’exemple externe 1. Ce test vérifie que les poids, les traitements et les fichiers fonctionnent ; il ne mesure pas combien de prédictions sont correctes.')

page('Architecture et organisation du code')
table(['Composant','Technologie','Rôle'],[('Interface','React TypeScript Vite','Écrans, formulaires, filtres et interactions.'),('Design','Tailwind CSS et CSS','Styles adaptés à la maquette, responsive.'),('API','Python FastAPI','Validation, comptes et accès aux fonctions.'),('Base','PostgreSQL SQLAlchemy Alembic','Métadonnées, analyses, décisions et migrations.'),('IA','YOLO Ultralytics PyTorch','Prédictions de segmentation.'),('Images','OpenCV NumPy Pillow','Lecture, orientation et composition des masques.'),('Carte','Leaflet OpenStreetMap','Affichage des coordonnées connues.')],[3.1,5.3,8.6])
h('Dossiers')
p('frontend/src contient les écrans React. App.tsx organise la navigation ; forms.tsx gère connexion, import et paramètres ; AnalysisDetail.tsx présente les résultats ; components.tsx regroupe tableau, carte et composants communs. api.ts centralise les appels au serveur.')
p('backend/app/main.py définit les routes. database.py définit les tables. auth.py gère les mots de passe et les sessions. segmentation.py charge le modèle et calcule les fichiers. jobs.py exécute les traitements. backend/migrations fait évoluer le schéma via Alembic.')
p('data/originals conserve les photographies et data/results les résultats par analyse. models conserve les poids et leur provenance. samples contient les exemples attribués. scripts contient installation, conversion, entraînement, évaluation et sauvegarde.')
h('Modes de traitement')
p('Sur Windows, la file est enregistrée en base et exécutée par un seul thread. Jusqu’à vingt analyses peuvent être en attente ; une observation ne peut pas posséder deux traitements simultanés. Un arrêt conserve les tâches en attente et signale les analyses interrompues lors du redémarrage.')
p('La variante compose.yaml fournit PostgreSQL, Redis, le serveur et un worker Celery. Docker n’étant pas installé sur ce PC, cette variante n’a pas été exécutée. Le mode local testé ne dépend pas de Redis. CVAT est un outil externe d’annotation, utilisé avec le convertisseur fourni.')
h('Accès')
p('Les mots de passe sont salés et hachés par scrypt. Les sessions sont stockées par empreinte de jeton et transmises par un cookie HttpOnly. Les rôles sont contrôlés côté API. Les fichiers d’image exigent une connexion, comme les données du registre.')

page('Préparer une expérimentation pour le mémoire')
h('Collecte et annotation')
p('Constituez un corpus de photographies de Goma avec déchets et sans déchets, dans plusieurs conditions d’éclairage, de distance et d’occultation. Consignez la provenance, le lieu réellement connu, la date et les conditions de capture. Les exemples externes fournis ne suffisent pas pour valider le terrain.')
p('Dans CVAT, annotez les zones de déchets avec des polygones et une nomenclature stable. Exportez les annotations au format COCO 1.0. Le convertisseur du projet accepte les polygones ; les annotations RLE nécessitent une conversion préalable et ne sont pas interprétées automatiquement.')
code('.venv\\Scripts\\python.exe scripts\\convert_coco.py\n  --annotations annotations.json --images images\n  --output datasets\\goma')
p('Le convertisseur crée train, val et test avec seed 42. Avant entraînement, regroupez les images du même site ou issues de la même capture dans une seule partition. Le contrôle automatique détecte les images identiques et les annotations absentes ; il ne détecte pas toutes les images presque identiques.')
h('Entraîner et activer')
code('.venv\\Scripts\\python.exe scripts\\train.py\n  --data datasets\\goma\\dataset.yaml --epochs 50 --device cpu')
p('Le terminal indique le dossier final des poids. Utilisez best.pt de ce dossier. Le calcul sur CPU peut être long ; un GPU exige une installation PyTorch CUDA compatible.')
code('.venv\\Scripts\\python.exe scripts\\register_model.py\n  --weights runs\\goma-segmentation\\weights\\best.pt\n  --name "YOLO déchets Goma" --dataset "Goma version 1"')
p('Redémarrez le serveur après activation du modèle. Ses poids et ses classes sont vérifiés avant utilisation.')
h('Évaluer sur le test indépendant')
code('.venv\\Scripts\\python.exe scripts\\evaluate.py\n  --data datasets\\goma\\dataset.yaml\n  --weights models\\goma-segmentation.pt --device cpu')
p('L’évaluation utilise la partition test, calcule la mAP des masques et sauvegarde les métriques avec l’empreinte des données. La page Modèle affiche l’évaluation uniquement lorsque son empreinte de poids correspond au modèle actif. Comparez les méthodes et les seuils avec un protocole stable ; ne choisissez pas les paramètres sur le test final.')

page('Installation sauvegarde et dépannage')
h('Installation et lancement')
p('Le PC dispose de Python 3.11, Node.js et PostgreSQL 18. Les dépendances et l’interface sont déjà installées dans PROJECT. PostgreSQL du projet utilise une instance dédiée sur 5433, distincte du service existant sur 5432. Les mots de passe de base sont générés dans les fichiers privés du projet.')
code('powershell -ExecutionPolicy Bypass -File .\\install.ps1\npowershell -ExecutionPolicy Bypass -File .\\start.ps1')
p('install.ps1 installe Python, compile React et télécharge les poids vérifiés. start.ps1 démarre PostgreSQL du projet, applique les migrations et lance FastAPI. Sur un autre PC, prévoyez Python 3.11, Node.js récent, PostgreSQL et l’espace disque nécessaire à PyTorch.')
h('Sauvegarde et restauration')
p('Attendez la fin des analyses, puis arrêtez l’API avec Ctrl+C. Laissez PostgreSQL actif. Exécutez la sauvegarde ; elle produit database.dump et files.zip dans un nouveau dossier daté de backups.')
code('.venv\\Scripts\\python.exe scripts\\backup.py')
p('Pour restaurer, utilisez pg_restore vers une nouvelle base PostgreSQL vide, puis extrayez files.zip dans une nouvelle copie du projet. Adaptez DATABASE_URL dans .env et démarrez cette copie. Conservez l’original jusqu’à vérification du nombre d’observations, des images et des résultats. Une restauration ne consiste pas à remplacer à chaud le dossier runtime/postgres.')
h('Problèmes courants')
table(['Symptôme','Action'],[('Adresse 8001 inaccessible','Lancez LANCER.cmd et consultez le terminal.'),('Port 8001 déjà occupé','Utilisez le serveur déjà lancé ou arrêtez son ancienne fenêtre.'),('Fond de carte absent','La carte OpenStreetMap nécessite Internet ; les images et analyses restent locales.'),('Analyse en échec','Consultez le journal, vérifiez les poids et relancez depuis l’observation.'),('Connexion oubliée','Utilisez scripts/reset_password.py --login IDENTIFIANT depuis le PC du projet.')],[6,11])
h('Commandes de contrôle')
code('.venv\\Scripts\\python.exe -m pytest tests -q\nnpm.cmd --prefix frontend run build\n.venv\\Scripts\\python.exe -m pip check')
p('Les tests automatisés contrôlent les droits et les entrées dans une base temporaire. Le test navigateur a aussi exécuté les vrais poids YOLO avec PostgreSQL.')

page('Périmètre académique et sources')
h('Correspondance avec le Word')
p('L’objectif 0.6.1 du document vise la détection, la segmentation et la localisation des déchets à partir d’images. La délimitation 0.7.2 retient les images et le contexte de Goma. Le logiciel suit ce périmètre avec une interface de consultation et de vérification pour l’aide à la décision.')
p('Le Word contient encore des éléments de gabarit de parking et de contrôle d’accès, une référence à Kinshasa et des sections répétées. Ses questions de recherche parlent de vidéo en temps réel alors que sa délimitation retient les images. Ces points doivent être harmonisés dans le mémoire avec l’encadreur.')
p('Un passage annonce également des cartes topologiques 3D et des invariants comme Euler et Betti. Aucun corpus volumique ni dispositif d’acquisition correspondant n’est fourni. Le périmètre 2D avec YOLO a été confirmé pour cette réalisation. Cette application n’implémente pas la méthode 3D et n’en revendique pas les garanties. Le passage 3D du mémoire doit être harmonisé avec ce périmètre.')
h('Travail expérimental restant')
p('Pour défendre le TFE, constituez et annotez un jeu indépendant de Goma, définissez les critères d’acceptation, mesurez les erreurs et comparez les variantes retenues. La réussite du lancement et de l’export des masques n’établit pas l’efficacité municipale ni une précision de détection. Le journal de vérification aide à examiner les résultats mais ne remplace pas le protocole d’évaluation.')
h('Sources et licences')
p('Modèle et métriques déclarées : https://huggingface.co/CatSat/yolov11-litter-materials')
p('Code et images exemples : https://github.com/CatSatOK/litter-detection-yolov11')
p('Documentation de segmentation : https://docs.ultralytics.com/tasks/segment/')
p('Annotation CVAT : https://docs.cvat.ai/')
p('Cartographie : https://www.openstreetmap.org/copyright')
p('Le fournisseur du modèle déclare MIT pour son dépôt. Le moteur Ultralytics utilise ses propres conditions AGPL-3.0 ou Enterprise. L’utilisation locale pour le TFE et une distribution commerciale constituent des contextes distincts ; examinez les licences des poids, du moteur et des images avant une distribution.')
p('La maquette utilisateur sert de référence visuelle. Ses chiffres et ses lieux ne constituent pas des observations enregistrées. Le fichier Word original est conservé sans modification dans REDACTION.')

doc.save(root/'docs/GUIDE_UTILISATION.docx')
print('GUIDE_UTILISATION.docx créé')
